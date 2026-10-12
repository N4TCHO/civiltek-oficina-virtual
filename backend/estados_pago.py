"""Estados de pago: cálculos auditables sin modificar el presupuesto aprobado."""
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from fastapi import HTTPException
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field

def number(value):
    try:
        n = Decimal(str(value))
        if not n.is_finite():
            raise ValueError()
        return n
    except (InvalidOperation, ValueError):
        raise HTTPException(422, "Valor numérico inválido")

def fmt(value):
    return float(value.quantize(Decimal("0.000001"), rounding=ROUND_HALF_UP))

class Item(BaseModel):
    codigo: str = Field(min_length=1, max_length=80)
    descripcion: str = Field(min_length=1, max_length=500)
    unidad: str = Field(default="un", max_length=30)
    cantidad_contrato: str
    precio_unitario: str
    acumulado_anterior: str = "0"
    cantidad_periodo: str = "0"

class EstadoPago(BaseModel):
    contrato: str = Field(min_length=1, max_length=300)
    numero: str = Field(min_length=1, max_length=30)
    periodo: str = Field(default="", max_length=100)
    items: list[Item] = Field(min_length=1, max_length=1500)

def calcular(ep):
    result=[]
    seen=set()
    for item in ep.items:
        if item.codigo in seen:
            raise HTTPException(422, "Código de partida duplicado: "+item.codigo)
        seen.add(item.codigo)
        q=number(item.cantidad_contrato); p=number(item.precio_unitario)
        anterior=number(item.acumulado_anterior); periodo=number(item.cantidad_periodo)
        if min(q,p,anterior,periodo)<0:
            raise HTTPException(422, "No se admiten valores negativos")
        total=anterior+periodo
        if total>q:
            raise HTTPException(422, "Cantidad acumulada supera el contrato en "+item.codigo)
        result.append({"codigo":item.codigo,"descripcion":item.descripcion,"unidad":item.unidad,
            "cantidad_contrato":fmt(q),"precio_unitario":fmt(p),"acumulado_anterior":fmt(anterior),
            "cantidad_periodo":fmt(periodo),"acumulado_actual":fmt(total),"saldo_cantidad":fmt(q-total),
            "monto_contrato":fmt(q*p),"monto_anterior":fmt(anterior*p),
            "monto_periodo":fmt(periodo*p),"monto_acumulado":fmt(total*p),
            "saldo_monto":fmt((q-total)*p),"avance_pct":fmt((total/q*100) if q else Decimal(0))})
    return result

def totales(rows):
    fields=("monto_contrato","monto_anterior","monto_periodo","monto_acumulado","saldo_monto")
    return {key:fmt(sum((number(row[key]) for row in rows),Decimal(0))) for key in fields}

def excel_ep(ep):
    rows=calcular(ep)
    wb=Workbook();ws=wb.active;ws.title="Estado de Pago"
    ws.append(["CIVILTEK | ESTADO DE PAGO"])
    ws.append(["Contrato",ep.contrato])
    ws.append(["Estado de pago",ep.numero,"Período",ep.periodo])
    ws.append(["Valores netos referenciales, sin reajustes, retenciones ni modificaciones contractuales"])
    ws.append([])
    headers=["Código","Descripción","Unidad","Cant. contrato","P. unitario","Acum. anterior","Cantidad período",
        "Acum. actual","Saldo cantidad","Monto contrato","Monto anterior","Monto período","Monto acumulado",
        "Saldo monto","Avance %"]
    ws.append(headers)
    keys=["codigo","descripcion","unidad","cantidad_contrato","precio_unitario","acumulado_anterior",
        "cantidad_periodo","acumulado_actual","saldo_cantidad","monto_contrato","monto_anterior",
        "monto_periodo","monto_acumulado","saldo_monto","avance_pct"]
    for row in rows:ws.append([row[k] for k in keys])
    summary=totales(rows);ws.append(["TOTAL","","","","","","","","",summary["monto_contrato"],
        summary["monto_anterior"],summary["monto_periodo"],summary["monto_acumulado"],summary["saldo_monto"]])
    ws.freeze_panes="D7";ws.auto_filter.ref=f"A6:O{6+len(rows)}"
    for col,width in {"A":17,"B":55,"C":12,"D":19,"E":18,"F":19,"G":19,"H":19,"I":19,
                      "J":22,"K":22,"L":22,"M":22,"N":22,"O":15}.items():
        ws.column_dimensions[col].width=width
    for cell in ws[6]:
        cell.fill=PatternFill("solid",fgColor="163047");cell.font=Font(color="FFFFFF",bold=True)
        cell.alignment=Alignment(wrap_text=True)
    for row in ws.iter_rows(min_row=7,min_col=4,max_col=15):
        for cell in row:cell.number_format='#,##0.00;[Red](#,##0.00)'
    ws["A1"].font=Font(size=16,bold=True,color="168F88")
    ws.print_options.horizontalCentered=True
    ws.sheet_properties.pageSetUpPr.fitToPage=True
    ws.page_setup.orientation="landscape";ws.page_setup.paperSize=ws.PAPERSIZE_A3
    output=BytesIO();wb.save(output);output.seek(0)
    return StreamingResponse(output,media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition":'attachment; filename="civiltek-estado-pago.xlsx"'})
