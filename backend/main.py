"""Civiltek API: single-admin MVP. Configure a private persistent volume before production."""
import hmac
import hashlib
import base64
import time
import io
import json
import os
import sqlite3
import threading
import uuid
from contextlib import contextmanager
from datetime import datetime, timezone
from pathlib import Path

import httpx
from docx import Document
from backend.estados_pago import EstadoPago, calcular, totales, excel_ep, importar_excel
from backend.operations import configure as configure_operations
from backend.licitaciones import configure as configure_licitaciones
from backend.apu import router as apu_router
from fastapi import FastAPI, Depends, Header, HTTPException, Request, Response, UploadFile, File
from fastapi.responses import FileResponse, StreamingResponse
from pydantic import BaseModel, Field

ROOT = Path(__file__).resolve().parent.parent
DB_PATH = Path(os.getenv("CIVILTEK_DB_PATH", str(ROOT / "data" / "civiltek.db")))
TOKEN = os.getenv("CIVILTEK_ADMIN_TOKEN", "")
LOCK = threading.RLock()
app = FastAPI(title="Civiltek API", version="0.4.0", docs_url=None, redoc_url=None, openapi_url=None)

ROLES = {
    "Director": "Coordina equipos multidisciplinarios, define alcance, prioridades y entregables.",
    "Licitaciones": "Analiza bases, requisitos de admisibilidad y propuestas de licitaciones chilenas.",
    "Presupuestos": "Elabora análisis de precios unitarios, cubicaciones y presupuestos verificables.",
    "Calidad MOP": "Desarrolla planes de calidad, protocolos, matrices y auditorías para contratos MOP.",
    "Documental": "Estructura expedientes, cartas, registros, procedimientos y anexos.",
    "Comercial": "Desarrolla ofertas de servicios, propuestas y comunicaciones comerciales.",
    "Auditor": "Revisa trazabilidad, coherencia, riesgos y evidencia de los entregables.",
    "Legal y Contratos": "Analiza contratos, controversias, obligaciones y normativa chilena; no sustituye asesoría ni patrocinio de abogado habilitado.",
    "Prevención de Riesgos": "Prepara matrices IPER, procedimientos seguros, planes de emergencia y programas preventivos sujetos a validación profesional.",
    "Medio Ambiente": "Prepara planes de manejo ambiental, matrices de aspectos e impactos, residuos y seguimiento de compromisos, sujetos a validación.",
    "Ingeniería y Obras": "Desarrolla memorias, especificaciones, cubicaciones y revisión técnica de proyectos; no certifica ni firma diseños.",
    "Finanzas": "Evalúa costos, flujo de caja, estados de pago y escenarios financieros.",
    "Topografía y GIS": "Revisa antecedentes topográficos, georreferenciación, cubicaciones y entregables cartográficos.",
    "Recursos Humanos": "Prepara políticas, procedimientos y documentación laboral con revisión normativa."
}

def utcnow():
    return datetime.now(timezone.utc).isoformat()

@contextmanager
def connect():
    DB_PATH.parent.mkdir(parents=True, exist_ok=True)
    with LOCK:
        db = sqlite3.connect(DB_PATH, timeout=20)
        db.row_factory = sqlite3.Row
        try:
            yield db
            db.commit()
        finally:
            db.close()

@app.on_event("startup")
def setup():
    with connect() as db:
        db.execute("""CREATE TABLE IF NOT EXISTS records(
            id TEXT PRIMARY KEY, kind TEXT NOT NULL, title TEXT NOT NULL,
            payload TEXT NOT NULL, created_at TEXT NOT NULL, updated_at TEXT NOT NULL
        )""")
        db.execute("""CREATE INDEX IF NOT EXISTS idx_records_kind ON records(kind)""")

SESSION_SECONDS = 7 * 24 * 60 * 60
COOKIE_NAME = "civiltek_session"

def _signature(value: str) -> str:
    digest = hmac.new(TOKEN.encode(), value.encode(), hashlib.sha256).digest()
    return base64.urlsafe_b64encode(digest).decode().rstrip("=")

def _valid_session(cookie: str | None) -> bool:
    if not cookie or "." not in cookie or len(TOKEN) < 32:
        return False
    timestamp, signature = cookie.split(".", 1)
    if not timestamp.isdigit() or not hmac.compare_digest(_signature(timestamp), signature):
        return False
    age = time.time() - int(timestamp)
    return 0 <= age <= SESSION_SECONDS

def authorize(request: Request, authorization: str | None = Header(default=None)):
    if len(TOKEN) < 32:
        raise HTTPException(503, "Configure CIVILTEK_ADMIN_TOKEN with at least 32 characters.")
    if authorization and hmac.compare_digest(authorization, "Bearer " + TOKEN):
        return
    if _valid_session(request.cookies.get(COOKIE_NAME)):
        if request.method not in ("GET", "HEAD", "OPTIONS"):
            origin = request.headers.get("origin", "")
            if origin != "https://" + request.headers.get("host", "") and origin != "http://" + request.headers.get("host", ""):
                raise HTTPException(403, "Origen de solicitud no permitido.")
        return
    raise HTTPException(401, "Inicia sesión para continuar.")

class LoginIn(BaseModel):
    password: str = Field(min_length=1, max_length=500)

@app.post("/api/auth/login")
def login(data: LoginIn, response: Response):
    if len(TOKEN) < 32:
        raise HTTPException(503, "Acceso no configurado.")
    if not hmac.compare_digest(data.password, TOKEN):
        raise HTTPException(401, "Clave incorrecta.")
    stamp = str(int(time.time()))
    response.set_cookie(COOKIE_NAME, stamp + "." + _signature(stamp),
        max_age=SESSION_SECONDS, httponly=True, secure=True, samesite="lax", path="/")
    response.headers["Cache-Control"] = "no-store"
    return {"ok": True, "expires_in_days": 7}

@app.post("/api/auth/logout")
def logout(response: Response):
    response.delete_cookie(COOKIE_NAME, path="/", secure=True, httponly=True, samesite="lax")
    return {"ok": True}

@app.get("/api/auth/status")
def auth_status(request: Request):
    return {"authenticated": _valid_session(request.cookies.get(COOKIE_NAME)), "session_days": 7}


class RecordIn(BaseModel):
    kind: str
    title: str = Field(min_length=1, max_length=200)
    payload: dict = Field(default_factory=dict)

class TaskIn(BaseModel):
    role: str
    instructions: str = Field(min_length=3, max_length=12000)
    context: str = Field(default="", max_length=30000)

@app.get("/api/health")
def health():
    return {"ok": True, "version": "0.4.0", "configured": len(TOKEN) >= 32}

@app.get("/api/records", dependencies=[Depends(authorize)])
def list_records(kind: str | None = None):
    with connect() as db:
        rows = db.execute("SELECT * FROM records WHERE (? IS NULL OR kind = ?) ORDER BY updated_at DESC LIMIT 500", (kind, kind)).fetchall()
    return [dict(r) | {"payload": json.loads(r["payload"])} for r in rows]

@app.post("/api/records", dependencies=[Depends(authorize)])
def create_record(record: RecordIn):
    if record.kind not in ("project", "task", "letter", "apu"):
        raise HTTPException(422, "Unsupported record kind")
    item_id = str(uuid.uuid4())
    now = utcnow()
    with connect() as db:
        db.execute("INSERT INTO records VALUES(?,?,?,?,?,?)", (item_id, record.kind, record.title, json.dumps(record.payload, ensure_ascii=False), now, now))
    return {"id": item_id, "created_at": now}

@app.put("/api/records/{record_id}", dependencies=[Depends(authorize)])
def update_record(record_id: str, record: RecordIn):
    if record.kind not in ("project", "task", "letter", "apu"):
        raise HTTPException(422, "Unsupported record kind")
    with connect() as db:
        row = db.execute("SELECT kind FROM records WHERE id=?", (record_id,)).fetchone()
        if not row:
            raise HTTPException(404, "Record not found")
        if row["kind"] != record.kind:
            raise HTTPException(409, "Cannot change record kind")
        db.execute("UPDATE records SET title=?, payload=?, updated_at=? WHERE id=?", (record.title, json.dumps(record.payload, ensure_ascii=False), utcnow(), record_id))
    return {"updated": True}

@app.delete("/api/records/{record_id}", dependencies=[Depends(authorize)])
def delete_record(record_id: str):
    with connect() as db:
        result = db.execute("DELETE FROM records WHERE id=?", (record_id,))
        if result.rowcount == 0:
            raise HTTPException(404, "Record not found")
    return {"deleted": True}

@app.post("/api/agents/run", dependencies=[Depends(authorize)])
async def run_agent(task: TaskIn):
    if task.role not in ROLES:
        raise HTTPException(422, "Unknown specialist")
    api_key = os.getenv("OPENROUTER_API_KEY", "")
    if not api_key:
        raise HTTPException(503, "OPENROUTER_API_KEY is not configured")
    model = os.getenv("OPENROUTER_MODEL", "openrouter/free")
    system = ("Eres un especialista de la consultora Civiltek, Chile. " + ROLES[task.role] +
        " Responde en español chileno profesional. No inventes leyes, valores de mercado, "
        "normas MOP, documentos ni resultados de ensayos. Identifica los datos faltantes, "
        "explicita supuestos y exige revisión humana antes de presentar documentos oficiales. "
        "No ejecutes acciones externas ni afirmes que las ejecutaste.")
    payload = {"model": model, "messages": [
        {"role": "system", "content": system},
        {"role": "user", "content": task.instructions + ("\n\nAntecedentes:\n" + task.context if task.context else "")}
    ], "temperature": 0.2, "max_tokens": 1800}
    # OpenRouter's free router selects an available free model. A paid model with
    # insufficient credits gets one automatic retry using that router.
    free_router = "openrouter/free"
    selected_model = model
    try:
        async with httpx.AsyncClient(timeout=75) as client:
            response = await client.post(
                "https://openrouter.ai/api/v1/chat/completions",
                headers={"Authorization": "Bearer " + api_key, "Content-Type": "application/json"},
                json=payload,
            )
            if response.status_code == 402 and selected_model != free_router:
                selected_model = free_router
                payload["model"] = free_router
                response = await client.post(
                    "https://openrouter.ai/api/v1/chat/completions",
                    headers={"Authorization": "Bearer " + api_key, "Content-Type": "application/json"},
                    json=payload,
                )
            if response.is_error:
                # Never return provider bodies: they can contain sensitive request data.
                status = response.status_code
                if status == 401 or status == 403:
                    detail = "OpenRouter rechazó la API Key o sus permisos (HTTP " + str(status) + ")."
                elif status == 402:
                    detail = "OpenRouter indica saldo insuficiente incluso para esta solicitud (HTTP 402)."
                elif status == 429:
                    detail = "Límite de solicitudes de OpenRouter alcanzado (HTTP 429). Intenta más tarde."
                elif status == 404:
                    detail = "Modelo de OpenRouter no disponible (HTTP 404)."
                else:
                    detail = "OpenRouter devolvió HTTP " + str(status) + "."
                raise HTTPException(502, detail)
            data = response.json()
            answer = data["choices"][0]["message"]["content"]
            if not isinstance(answer, str) or not answer.strip():
                raise ValueError("Empty model answer")
    except (httpx.HTTPError, KeyError, IndexError, ValueError, TypeError) as exc:
        raise HTTPException(502, "No fue posible obtener respuesta de OpenRouter. Comprueba disponibilidad del proveedor.") from exc
    return {"role": task.role, "model": selected_model, "draft": answer, "requires_human_review": True}

# Official-source opportunity intake: only public procurement metadata.
# The API ticket stays on the server; never return it to the browser.
class OpportunityScanIn(BaseModel):
    terms: list[str] = Field(default_factory=lambda: [
        "asesoría", "prevención de riesgos", "medio ambiente", "auditoría",
        "ingeniería", "topografía", "conservación vial", "consultoría"
    ], max_length=20)

@app.get("/api/integrations/status", dependencies=[Depends(authorize)])
def integration_status():
    return {"mercado_publico": {"configured": bool(os.getenv("MERCADOPUBLICO_TICKET")),
                                 "source": "ChileCompra API oficial",
                                 "requires_ticket": True},
            "economic": {"source": "Banco Central BDE", "configured": False},
            "environment": {"source": "IDE MMA WMS/WFS", "configured": False}}

@app.post("/api/opportunities/scan", dependencies=[Depends(authorize)])
async def scan_opportunities(scan: OpportunityScanIn):
    ticket = os.getenv("MERCADOPUBLICO_TICKET", "").strip()
    if not ticket:
        raise HTTPException(503, "Falta MERCADOPUBLICO_TICKET en Railway. No se ha consultado ChileCompra.")
    terms = [term.strip().lower() for term in scan.terms if 2 <= len(term.strip()) <= 100]
    if not terms:
        raise HTTPException(422, "Indica al menos una palabra clave válida.")
    url = "https://api.mercadopublico.cl/servicios/v1/publico/licitaciones.json"
    try:
        async with httpx.AsyncClient(timeout=35) as client:
            response = await client.get(url, params={"estado": "activas", "ticket": ticket})
            response.raise_for_status()
            data = response.json()
    except (httpx.HTTPError, ValueError) as exc:
        raise HTTPException(502, "No se pudo consultar la API oficial de ChileCompra.") from exc
    if not isinstance(data, dict) or not isinstance(data.get("Listado"), list):
        raise HTTPException(502, "ChileCompra devolvió una estructura inesperada.")
    matches = []
    for item in data["Listado"]:
        if not isinstance(item, dict):
            continue
        code = str(item.get("CodigoExterno", "")).strip()
        name = str(item.get("Nombre", "")).strip()
        description = str(item.get("Descripcion", "")).strip()
        if not code or not name or not any(t in (name + " " + description).lower() for t in terms):
            continue
        matches.append({"code": code, "name": name, "description": description[:1200],
                        "close_date": item.get("FechaCierre"),
                        "source": "Dirección ChileCompra", "status": "Por evaluar"})
    saved = 0
    with connect() as db:
        for item in matches[:200]:
            existing = db.execute("SELECT id FROM records WHERE kind='opportunity' AND json_extract(payload, '$.code')=?",
                                  (item["code"],)).fetchone()
            if existing:
                continue
            now = utcnow()
            db.execute("INSERT INTO records VALUES(?,?,?,?,?,?)",
                       (str(uuid.uuid4()), "opportunity", item["name"][:200],
                        json.dumps(item, ensure_ascii=False), now, now))
            saved += 1
    return {"source": "Dirección ChileCompra", "matches": len(matches),
            "new_saved": saved, "results": matches[:100],
            "note": "Coincidencia por palabras clave, no evaluación de admisibilidad."}

@app.get("/api/opportunities", dependencies=[Depends(authorize)])
def list_opportunities():
    with connect() as db:
        rows = db.execute("SELECT * FROM records WHERE kind='opportunity' ORDER BY created_at DESC LIMIT 200").fetchall()
    return [dict(r) | {"payload": json.loads(r["payload"])} for r in rows]


class ClientIn(BaseModel):
    name: str = Field(min_length=2, max_length=200)
    contact: str = Field(default="", max_length=250)
    notes: str = Field(default="", max_length=2000)

class ClientRequestIn(BaseModel):
    client_id: str
    service: str = Field(min_length=3, max_length=200)
    objective: str = Field(min_length=12, max_length=8000)
    context: str = Field(default="", max_length=16000)
    deadline: str = Field(default="", max_length=30)

@app.get("/api/clients", dependencies=[Depends(authorize)])
def list_clients():
    with connect() as db:
        rows = db.execute("SELECT * FROM records WHERE kind='client' ORDER BY updated_at DESC LIMIT 500").fetchall()
    return [dict(r) | {"payload": json.loads(r["payload"])} for r in rows]

@app.post("/api/clients", dependencies=[Depends(authorize)])
def add_client(client: ClientIn):
    item_id, now = str(uuid.uuid4()), utcnow()
    payload = {"name": client.name, "contact": client.contact, "notes": client.notes}
    with connect() as db:
        db.execute("INSERT INTO records VALUES(?,?,?,?,?,?)",
                   (item_id, "client", client.name, json.dumps(payload, ensure_ascii=False), now, now))
    return {"id": item_id, "name": client.name}

@app.get("/api/client-requests", dependencies=[Depends(authorize)])
def list_client_requests():
    with connect() as db:
        rows = db.execute("SELECT * FROM records WHERE kind='client_request' ORDER BY updated_at DESC LIMIT 500").fetchall()
    return [dict(r) | {"payload": json.loads(r["payload"])} for r in rows]

@app.post("/api/client-requests", dependencies=[Depends(authorize)])
def add_client_request(item: ClientRequestIn):
    with connect() as db:
        client = db.execute("SELECT id,title FROM records WHERE id=? AND kind='client'", (item.client_id,)).fetchone()
        if not client:
            raise HTTPException(404, "Cliente no encontrado")
        item_id, now = str(uuid.uuid4()), utcnow()
        payload = item.model_dump() if hasattr(item, "model_dump") else item.dict()
        payload.update({"client_name": client["title"], "status": "Pendiente de asignación", "office_id": None})
        db.execute("INSERT INTO records VALUES(?,?,?,?,?,?)",
                   (item_id, "client_request", item.service + ": " + item.objective[:150],
                    json.dumps(payload, ensure_ascii=False), now, now))
    return {"id": item_id, "status": "Pendiente de asignación"}

@app.post("/api/client-requests/{request_id}/dispatch", dependencies=[Depends(authorize)])
async def dispatch_client_request(request_id: str):
    with connect() as db:
        row = db.execute("SELECT payload FROM records WHERE id=? AND kind='client_request'", (request_id,)).fetchone()
    if not row:
        raise HTTPException(404, "Solicitud no encontrada")
    payload = json.loads(row["payload"])
    if payload.get("office_id"):
        raise HTTPException(409, "Esta solicitud ya tiene un expediente")
    result = await execute_office(OfficeRequest(
        objective=payload["service"] + ": " + payload["objective"],
        project=payload["client_name"],
        context=payload.get("context", "") + "\nPlazo indicado: " + payload.get("deadline", "")))
    with connect() as db:
        latest = db.execute("SELECT payload FROM records WHERE id=? AND kind='client_request'", (request_id,)).fetchone()
        latest_payload = json.loads(latest["payload"])
        latest_payload.update({"status": "Revisión humana pendiente", "office_id": result["id"]})
        db.execute("UPDATE records SET payload=?,updated_at=? WHERE id=?",
                   (json.dumps(latest_payload, ensure_ascii=False), utcnow(), request_id))
    return {"id": request_id, "office_id": result["id"], "status": "Revisión humana pendiente"}


class OfficeRequest(BaseModel):
    objective: str = Field(min_length=12, max_length=8000)
    project: str = Field(default="Civiltek", max_length=200)
    context: str = Field(default="", max_length=16000)

def _agent_system(role: str) -> str:
    return ("Eres el agente " + role + " de Civiltek, consultora chilena de construcción. "
            + ROLES[role] + " Trabaja en español profesional. "
            "No inventes datos, precios, normas, numerales MOP, documentos, ensayos ni leyes. "
            "Separa antecedentes verificados, supuestos, vacíos y propuestas. "
            "Todo entregable debe revisarse por un profesional antes de uso oficial.")

async def _ask_provider(client: httpx.AsyncClient, role: str, instruction: str, context: str) -> dict:
    api_key = os.getenv("OPENROUTER_API_KEY", "")
    if not api_key:
        raise HTTPException(503, "Configura OPENROUTER_API_KEY en Railway.")
    model = os.getenv("OPENROUTER_MODEL", "openrouter/free")
    messages = [{"role": "system", "content": _agent_system(role)},
                {"role": "user", "content": instruction + "\n\nANTECEDENTES:\n" + context}]
    payload = {"model": model, "messages": messages, "max_tokens": 1250, "temperature": 0.2}
    headers = {"Authorization": "Bearer " + api_key, "Content-Type": "application/json"}
    for attempt in range(2):
        try:
            response = await client.post("https://openrouter.ai/api/v1/chat/completions",
                                         json=payload, headers=headers)
        except httpx.HTTPError as exc:
            raise HTTPException(502, "No fue posible contactar OpenRouter.") from exc
        if response.status_code == 402 and model != "openrouter/free" and attempt == 0:
            model = "openrouter/free"
            payload["model"] = model
            continue
        if response.is_error:
            raise HTTPException(502, "OpenRouter rechazó la solicitud (HTTP "
                                + str(response.status_code) + ").")
        try:
            content = response.json()["choices"][0]["message"]["content"]
            if not isinstance(content, str) or not content.strip():
                raise ValueError("Respuesta vacía")
            return {"role": role, "model": model, "text": content}
        except (ValueError, KeyError, IndexError, TypeError) as exc:
            raise HTTPException(502, "El modelo no entregó texto válido.") from exc
    raise HTTPException(502, "No fue posible usar el modelo.")

@app.get("/api/office/agents", dependencies=[Depends(authorize)])
def office_agent_directory():
    """Real available role definitions; readiness is not a background-running claim."""
    return {"agents":[{"name":name,"description":description,
                        "category":"Dirección" if name=="Director" else
                                   "Control" if name=="Auditor" else "Especialista"}
                       for name,description in ROLES.items()],
            "provider_configured":bool(os.getenv("OPENROUTER_API_KEY"))}

@app.get("/api/office/deliveries", dependencies=[Depends(authorize)])
def office_deliveries():
    with connect() as db:
        rows=db.execute("SELECT id,title,payload,created_at,updated_at FROM records WHERE kind='task' ORDER BY updated_at DESC LIMIT 40").fetchall()
    items=[]
    for row in rows:
        try:
            data=json.loads(row["payload"])
        except (ValueError,TypeError):
            continue
        items.append({"id":row["id"],"title":row["title"],
                      "project":data.get("project",""),"status":data.get("status","Sin estado"),
                      "specialists":[x.get("role","") for x in data.get("specialists",[])],
                      "created_at":row["created_at"],"updated_at":row["updated_at"]})
    return {"items":items}

@app.post("/api/office/execute", dependencies=[Depends(authorize)])
async def execute_office(request: OfficeRequest):
    """Director plans, specialists draft, auditor checks; save the expediente."""
    with httpx.AsyncClient(timeout=90) as client:
        director = await _ask_provider(
            client, "Director",
            "Clasifica el encargo y redacta un plan breve. Elige hasta DOS especialistas "
            "de esta lista: " + ", ".join(r for r in ROLES if r not in ("Director", "Auditor")) + ". "
            "Responde PRIMERO con una línea estricta ROLES: nombre1, nombre2 "
            "y después PLAN: y el plan. No uses otros roles.",
            "Proyecto: " + request.project + "\nObjetivo: " + request.objective
            + "\n" + request.context)
        import re
        allowed = tuple(role for role in ROLES if role not in ("Director", "Auditor"))
        match = re.search(r"(?im)^ROLES:\\s*(.+)$", director["text"])
        selected = []
        if match:
            selected = [r for r in allowed if r.lower() in match.group(1).lower()][:2]
        if not selected:
            selected = ["Documental"]
        outputs = []
        for role in selected:
            output = await _ask_provider(
                client, role,
                "Realiza TU parte del encargo y produce un entregable útil, concreto y "
                "estructurado. No digas que ejecutaste acciones externas. "
                "Si faltan documentos, deja pendientes específicos.",
                "Proyecto: " + request.project + "\nObjetivo: " + request.objective
                + "\nAntecedentes: " + request.context
                + "\nPlan del director: " + director["text"][:2500])
            outputs.append(output)
        auditor = await _ask_provider(
            client, "Auditor",
            "Audita los entregables: detecta afirmaciones no verificadas, riesgos, "
            "errores, datos faltantes y condiciones de aprobación. "
            "Emite un dictamen BORRADOR REQUIERE REVISION HUMANA.",
            "Objetivo: " + request.objective + "\n"
            + "\n".join(x["role"] + ": " + x["text"][:5000] for x in outputs))
    now = utcnow()
    item_id = str(uuid.uuid4())
    result = {"id": item_id, "project": request.project, "objective": request.objective,
              "director": director, "specialists": outputs, "auditor": auditor,
              "status": "Revisión humana pendiente", "created_at": now}
    with connect() as db:
        db.execute("INSERT INTO records VALUES(?,?,?,?,?,?)",
                   (item_id, "task", request.objective[:200],
                    json.dumps(result, ensure_ascii=False), now, now))
    return result

@app.post("/api/licitaciones/expedientes/{ident}/ejecutar", dependencies=[Depends(authorize)])
async def ejecutar_estudio_licitacion(ident: str):
    """Orquestación real sobre los antecedentes guardados, con revisión humana obligatoria."""
    with connect() as db:
        row = db.execute("SELECT payload FROM records WHERE id=? AND kind='tender'", (ident,)).fetchone()
    if not row:
        raise HTTPException(404, "Expediente no encontrado")
    tender = json.loads(row["payload"])
    docs = tender.get("documents", [])
    if not docs:
        raise HTTPException(422, "Primero incorpora las bases y anexos del expediente.")
    context = "LICITACIÓN: " + tender["name"] + "\nMANDANTE: " + tender.get("client", "")
    context += "\nDOCUMENTOS Y TEXTO EXTRAÍDO:\n"
    for doc in docs[:15]:
        context += "\n### FUENTE: " + doc["filename"] + " (id " + doc["id"] + ")\n"
        context += doc["text"][:12000] + "\n"
        if len(context) > 55000:
            break
    objective = ("Estudiar las bases de esta licitación y elaborar un expediente preliminar. "
                 "Identificar requisitos con citas al nombre del archivo y fragmentos literales, "
                 "requisitos de admisibilidad, garantías, plazos, antecedentes faltantes, "
                 "partidas y APU necesarios, metodología, programación y planes exigidos. "
                 "NO inventar precios, rendimientos, cubicaciones, normas ni condiciones.")
    return await execute_office(OfficeRequest(objective=objective,
                                                project=tender["name"], context=context))

@app.get("/api/office/{item_id}/docx", dependencies=[Depends(authorize)])
def office_docx(item_id: str):
    with connect() as db:
        row = db.execute("SELECT payload FROM records WHERE id=? AND kind='task'",
                         (item_id,)).fetchone()
    if not row:
        raise HTTPException(404, "Expediente no encontrado.")
    result = json.loads(row["payload"])
    if "director" not in result:
        raise HTTPException(422, "Este registro no es un expediente de oficina.")
    doc = Document()
    doc.add_heading("CIVILTEK | Expediente de trabajo", 0)
    doc.add_paragraph("Proyecto: " + result["project"])
    doc.add_paragraph("Encargo: " + result["objective"])
    doc.add_paragraph("Estado: REQUIERE REVISIÓN HUMANA")
    for title, body in [("Director | Plan", result["director"]["text"])] + [
        (x["role"] + " | Entregable", x["text"]) for x in result["specialists"]
    ] + [("Auditor | Observaciones", result["auditor"]["text"])]:
        doc.add_heading(title, level=1)
        for line in body.split("\n"):
            doc.add_paragraph(line)
    stream = io.BytesIO()
    doc.save(stream)
    stream.seek(0)
    return StreamingResponse(stream,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": 'attachment; filename="civiltek-expediente.docx"'})


class LetterIn(BaseModel):
    number: str = Field(max_length=50)
    date: str = Field(max_length=30)
    recipient: str = Field(max_length=250)
    subject: str = Field(max_length=300)
    body: str = Field(max_length=20000)
    signer: str = Field(max_length=250)

@app.post("/api/letters/docx", dependencies=[Depends(authorize)])
def letter_docx(letter: LetterIn):
    doc = Document()
    doc.add_heading("CARTA N.º " + letter.number, level=1)
    doc.add_paragraph(letter.date)
    doc.add_paragraph("A: " + letter.recipient)
    doc.add_paragraph("MAT.: " + letter.subject)
    for paragraph in letter.body.split("\n"):
        doc.add_paragraph(paragraph)
    doc.add_paragraph("Atentamente,")
    doc.add_paragraph(letter.signer)
    stream = io.BytesIO()
    doc.save(stream)
    stream.seek(0)
    return StreamingResponse(stream, media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": 'attachment; filename="civiltek-carta.docx"'})

@app.post("/api/estados-pago/importar", dependencies=[Depends(authorize)])
async def importar_estado_pago(file: UploadFile = File(...)):
    if not (file.filename or "").lower().endswith(".xlsx"):
        raise HTTPException(422, "Solo se aceptan archivos .xlsx")
    data = await file.read(5*1024*1024+1)
    return {"items": importar_excel(data)}

@app.get("/api/estados-pago/guardados", dependencies=[Depends(authorize)])
def listar_estados_pago():
    with connect() as db:
        rows = db.execute("SELECT id, title, payload, updated_at FROM records WHERE kind='payment' ORDER BY updated_at DESC LIMIT 100").fetchall()
    return [{"id": r["id"], "title": r["title"], "updated_at": r["updated_at"],
             "data": json.loads(r["payload"])} for r in rows]

@app.post("/api/estados-pago/guardar", dependencies=[Depends(authorize)])
def guardar_estado_pago(data: EstadoPago):
    calcular(data)
    payment_id = str(uuid.uuid4())
    with connect() as db:
        db.execute("INSERT INTO records VALUES (?,?,?,?,?,?)",
            (payment_id, "payment", data.contrato+" · EP "+data.numero,
             data.model_dump_json(), utcnow(), utcnow()))
    return {"id": payment_id, "saved": True}

@app.post("/api/estados-pago/calcular", dependencies=[Depends(authorize)])
def calcular_estado_pago(data: EstadoPago):
    rows = calcular(data)
    return {"contrato": data.contrato, "numero": data.numero, "periodo": data.periodo,
            "items": rows, "totales": totales(rows),
            "advertencia": "Cálculo base sin reajustes, retenciones, anticipos, IVA ni modificaciones contractuales."}

@app.post("/api/estados-pago/excel", dependencies=[Depends(authorize)])
def descargar_estado_pago(data: EstadoPago):
    return excel_ep(data)

@app.get("/centro-mando")
def centro_mando_page():
    return FileResponse(ROOT / "centro-mando.html")

@app.get("/apu")
def apu_page():
    return FileResponse(ROOT / "apu.html")

@app.get("/estudio-licitaciones")
def estudio_licitaciones_page():
    return FileResponse(ROOT / "estudio-licitaciones.html")

@app.get("/operaciones")
def operaciones_page():
    return FileResponse(ROOT / "operaciones.html")

@app.get("/estados-pago")
def estados_pago_page():
    return FileResponse(ROOT / "estados-pago.html")

@app.get("/civiltek-ui.css")
def civiltek_stylesheet():
    return FileResponse(ROOT / "civiltek-ui.css", media_type="text/css", headers={"Cache-Control": "public, max-age=300"})

@app.get("/")
def index():
    return FileResponse(ROOT / "index.html")

@app.get("/cartas")
def letters_page():
    return FileResponse(ROOT / "cartas.html")

@app.get("/agentes")
def agents_page():
    return FileResponse(ROOT / "agentes.html")

@app.get("/workspace")
def workspace_page():
    return FileResponse(ROOT / "workspace.html")

@app.get("/oficina")
def office_page():
    return FileResponse(ROOT / "centro-mando.html")

@app.get("/clientes")
def clients_page():
    return FileResponse(ROOT / "clientes.html")

@app.get("/oportunidades")
def opportunities_page():
    return FileResponse(ROOT / "oportunidades.html")

@app.get("/api/dashboard", dependencies=[Depends(authorize)])
def dashboard():
    with connect() as db:
        rows = db.execute("SELECT kind, payload, created_at, updated_at, id, title FROM records WHERE kind IN ('client','client_request','opportunity','task') ORDER BY updated_at DESC LIMIT 500").fetchall()
    clients, requests, opportunities, tasks = [], [], [], []
    for row in rows:
        try:
            payload = json.loads(row["payload"])
        except (ValueError, TypeError):
            continue
        item = {"id": row["id"], "title": row["title"], "created_at": row["created_at"],
                "updated_at": row["updated_at"], "payload": payload}
        {"client": clients, "client_request": requests, "opportunity": opportunities, "task": tasks}[row["kind"]].append(item)
    pending = [x for x in requests if x["payload"].get("status") != "Aprobado"]
    reviews = [x for x in requests if x["payload"].get("status") == "Revisión humana pendiente"]
    return {"counts": {"clients": len(clients), "requests": len(requests),
                       "pending": len(pending), "reviews": len(reviews),
                       "opportunities": len(opportunities), "dossiers": len(tasks)},
            "requests": requests[:40], "opportunities": opportunities[:15],
            "integration": {"mercado_publico": bool(os.getenv("MERCADOPUBLICO_TICKET")),
                            "openrouter": bool(os.getenv("OPENROUTER_API_KEY"))},
            "generated_at": utcnow()}

class ReviewDecision(BaseModel):
    decision: str
    notes: str = Field(default="", max_length=4000)

@app.post("/api/client-requests/{request_id}/review", dependencies=[Depends(authorize)])
def review_client_request(request_id: str, decision: ReviewDecision):
    if decision.decision not in ("approve", "changes"):
        raise HTTPException(422, "Decisión inválida")
    with connect() as db:
        row = db.execute("SELECT payload FROM records WHERE id=? AND kind='client_request'", (request_id,)).fetchone()
        if not row:
            raise HTTPException(404, "Encargo no encontrado")
        payload = json.loads(row["payload"])
        if payload.get("status") != "Revisión humana pendiente" or not payload.get("office_id"):
            raise HTTPException(409, "No hay expediente pendiente de revisión")
        payload["status"] = "Aprobado" if decision.decision == "approve" else "Requiere correcciones"
        payload["review_notes"] = decision.notes
        payload["reviewed_at"] = utcnow()
        db.execute("UPDATE records SET payload=?, updated_at=? WHERE id=?",
                   (json.dumps(payload, ensure_ascii=False), utcnow(), request_id))
    return {"id": request_id, "status": payload["status"],
            "note": "La aprobación interna no envía ni firma documentos."}

@app.get("/panel")
def dashboard_page():
    return FileResponse(ROOT / "panel.html")

# Documents are kept outside the public web root, in the same persistent volume as SQLite.
UPLOAD_DIR = DB_PATH.parent / "attachments"
ALLOWED_EXTENSIONS = {".pdf", ".docx", ".xlsx", ".csv", ".txt", ".png", ".jpg", ".jpeg"}
MAX_ATTACHMENT_BYTES = 12 * 1024 * 1024

@app.post("/api/client-requests/{request_id}/attachments", dependencies=[Depends(authorize)])
async def upload_attachment(request_id: str, file: UploadFile = File(...)):
    with connect() as db:
        parent = db.execute("SELECT id FROM records WHERE id=? AND kind='client_request'", (request_id,)).fetchone()
    if not parent:
        raise HTTPException(404, "Encargo no encontrado")
    original = Path(file.filename or "").name
    ext = Path(original).suffix.lower()
    if ext not in ALLOWED_EXTENSIONS or not original or len(original) > 180:
        raise HTTPException(422, "Tipo de archivo no admitido")
    content = await file.read(MAX_ATTACHMENT_BYTES + 1)
    await file.close()
    if len(content) > MAX_ATTACHMENT_BYTES:
        raise HTTPException(413, "El archivo supera los 12 MB")
    # Reject obvious mismatched file signatures for common binary formats.
    if ext == ".pdf" and not content.startswith(b"%PDF-"):
        raise HTTPException(422, "El contenido no corresponde a un PDF")
    if ext in {".docx", ".xlsx"} and not content.startswith(b"PK"):
        raise HTTPException(422, "Archivo Office inválido")
    item_id = str(uuid.uuid4())
    UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
    destination = UPLOAD_DIR / (item_id + ext)
    try:
        with destination.open("xb") as out:
            out.write(content)
        now = utcnow()
        metadata = {"request_id": request_id, "filename": original, "bytes": len(content),
                    "stored_name": destination.name, "sha256": hashlib.sha256(content).hexdigest()}
        with connect() as db:
            db.execute("INSERT INTO records VALUES(?,?,?,?,?,?)",
                       (item_id, "attachment", original,
                        json.dumps(metadata, ensure_ascii=False), now, now))
    except Exception:
        destination.unlink(missing_ok=True)
        raise
    return {"id": item_id, "filename": original, "bytes": len(content)}

@app.get("/api/client-requests/{request_id}/attachments", dependencies=[Depends(authorize)])
def list_attachments(request_id: str):
    with connect() as db:
        parent = db.execute("SELECT id FROM records WHERE id=? AND kind='client_request'", (request_id,)).fetchone()
        if not parent:
            raise HTTPException(404, "Encargo no encontrado")
        rows = db.execute("SELECT id,payload,created_at FROM records WHERE kind='attachment' AND json_extract(payload,'$.request_id')=? ORDER BY created_at DESC",
                          (request_id,)).fetchall()
    return [{"id": row["id"], "created_at": row["created_at"],
             "filename": json.loads(row["payload"])["filename"],
             "bytes": json.loads(row["payload"])["bytes"]} for row in rows]

@app.get("/api/client-requests/{request_id}/attachments/{attachment_id}", dependencies=[Depends(authorize)])
def download_attachment(request_id: str, attachment_id: str):
    with connect() as db:
        row = db.execute("SELECT payload FROM records WHERE id=? AND kind='attachment'", (attachment_id,)).fetchone()
    if not row:
        raise HTTPException(404, "Adjunto no encontrado")
    payload = json.loads(row["payload"])
    if payload["request_id"] != request_id:
        raise HTTPException(404, "Adjunto no encontrado")
    stored = payload["stored_name"]
    if Path(stored).name != stored:
        raise HTTPException(500, "Nombre de archivo inválido")
    path = UPLOAD_DIR / stored
    if not path.is_file():
        raise HTTPException(404, "Archivo no disponible en el volumen")
    return FileResponse(path, filename=payload["filename"],
                        media_type="application/octet-stream",
                        headers={"Cache-Control": "no-store", "X-Content-Type-Options": "nosniff"})

app.include_router(configure_operations(connect, utcnow), dependencies=[Depends(authorize)])

app.include_router(configure_licitaciones(connect, utcnow), dependencies=[Depends(authorize)])

app.include_router(apu_router, dependencies=[Depends(authorize)])
