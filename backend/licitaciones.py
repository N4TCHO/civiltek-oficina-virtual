"""Tender dossier intake: deterministic document extraction, never invent base requirements."""
import io
import json
import uuid
from datetime import datetime
from pathlib import Path
from fastapi import APIRouter, UploadFile, File, HTTPException
from pydantic import BaseModel, Field
from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

router=APIRouter(prefix="/api/licitaciones",tags=["licitaciones"])
MAX_BYTES=12*1024*1024
MAX_TEXT=120000
class TenderIn(BaseModel):
    name:str=Field(min_length=5,max_length=300)
    code:str=Field(default="",max_length=100)
    client:str=Field(default="",max_length=200)
    deadline:str=Field(default="",max_length=50)
    notes:str=Field(default="",max_length=4000)

def extract(filename,raw):
    ext=Path(filename).suffix.lower()
    if len(raw)>MAX_BYTES:raise HTTPException(413,"Archivo demasiado grande (máximo 12 MB)")
    try:
        if ext==".pdf":
            reader=PdfReader(io.BytesIO(raw))
            if reader.is_encrypted:raise HTTPException(422,"PDF protegido; requiere copia sin contraseña")
            text="\n".join((p.extract_text() or "") for p in reader.pages[:120])
        elif ext==".docx":
            doc=Document(io.BytesIO(raw))
            text="\n".join([p.text for p in doc.paragraphs]+[" | ".join(c.text for c in row.cells) for t in doc.tables for row in t.rows])
        elif ext==".xlsx":
            wb=load_workbook(io.BytesIO(raw),read_only=True,data_only=True)
            text="\n".join(f"[{ws.title}] "+ " | ".join(str(v) for v in row if v is not None)
                for ws in wb.worksheets[:8] for row in ws.iter_rows(max_row=1500,max_col=25,values_only=True))
            wb.close()
        elif ext==".txt":
            text=raw.decode("utf-8-sig")
        else:raise HTTPException(422,"Formatos permitidos: PDF, DOCX, XLSX, TXT")
    except HTTPException:raise
    except Exception as exc:raise HTTPException(422,"No fue posible extraer texto del archivo") from exc
    if not text.strip():raise HTTPException(422,"No se detectó texto; un PDF escaneado requiere OCR")
    return text[:MAX_TEXT],len(text)>MAX_TEXT

def configure(connect,utcnow):
    @router.post("/expedientes")
    def create(data:TenderIn):
        ident=str(uuid.uuid4());now=utcnow()
        payload=data.model_dump()|{"documents":[],"status":"Antecedentes en recopilación"}
        with connect() as db:
            db.execute("INSERT INTO records VALUES(?,?,?,?,?,?)",
                (ident,"tender",data.name,json.dumps(payload,ensure_ascii=False),now,now))
        return {"id":ident,"status":payload["status"]}

    @router.get("/expedientes")
    def list_all():
        with connect() as db:
            rows=db.execute("SELECT id,title,payload,updated_at FROM records WHERE kind='tender' ORDER BY updated_at DESC LIMIT 100").fetchall()
        return [{"id":r["id"],"title":r["title"],"updated_at":r["updated_at"],**json.loads(r["payload"])} for r in rows]

    @router.post("/expedientes/{ident}/documentos")
    async def upload(ident:str,file:UploadFile=File(...)):
        name=Path(file.filename or "").name
        raw=await file.read(MAX_BYTES+1)
        text,truncated=extract(name,raw)
        with connect() as db:
            row=db.execute("SELECT payload FROM records WHERE id=? AND kind='tender'",(ident,)).fetchone()
            if not row:raise HTTPException(404,"Expediente no encontrado")
            payload=json.loads(row["payload"])
            doc={"id":str(uuid.uuid4()),"filename":name,"characters":len(text),"truncated":truncated,
                 "text":text,"uploaded_at":utcnow()}
            payload["documents"].append(doc)
            if len(payload["documents"])>30:raise HTTPException(422,"Máximo 30 documentos por expediente")
            db.execute("UPDATE records SET payload=?,updated_at=? WHERE id=?",
                (json.dumps(payload,ensure_ascii=False),utcnow(),ident))
        return {"id":doc["id"],"filename":name,"characters":len(text),"truncated":truncated}

    @router.get("/expedientes/{ident}/matriz")
    def matrix(ident:str):
        with connect() as db:
            row=db.execute("SELECT payload FROM records WHERE id=? AND kind='tender'",(ident,)).fetchone()
        if not row:raise HTTPException(404,"Expediente no encontrado")
        payload=json.loads(row["payload"])
        sections=["Admisibilidad","Experiencia y registro","Garantías","Plazos","Presupuesto y APU",
                  "Plan de calidad","Prevención de riesgos","Medio ambiente","Programación","Anexos y formularios"]
        return {"expediente":ident,"documents":len(payload["documents"]),
                "checks":[{"area":s,"status":"Pendiente de revisión","source":None} for s in sections],
                "warning":"Matriz inicial de verificación. No implica que estos requisitos figuren en las bases; deben comprobarse con cita de fuente."}
    return router
