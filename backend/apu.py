"""Deterministic unit-price analysis; prices and yields must come from user inputs."""
from decimal import Decimal,InvalidOperation,ROUND_HALF_UP
from io import BytesIO
from fastapi import APIRouter,HTTPException
from fastapi.responses import StreamingResponse
from pydantic import BaseModel,Field
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill

router=APIRouter(prefix="/api/apu",tags=["apu"])
class Resource(BaseModel):
    name:str=Field(min_length=1,max_length=200)
    type:str
    unit:str=Field(min_length=1,max_length=25)
    quantity:str
    unit_price:str
class APU(BaseModel):
    code:str=Field(min_length=1,max_length=70)
    description:str=Field(min_length=1,max_length=400)
    unit:str=Field(min_length=1,max_length=25)
    resources:list[Resource]=Field(min_length=1,max_length=300)
    indirect_pct:str="0"
    profit_pct:str="0"
def dec(x):
    try:
        d=Decimal(str(x))
        if not d.is_finite() or d<0:raise ValueError()
        return d
    except (InvalidOperation,ValueError):
        raise HTTPException(422,"Cantidad, precio o porcentaje inválido")
def round2(x):return float(x.quantize(Decimal("0.01"),rounding=ROUND_HALF_UP))
def compute(data):
    amounts={"Materiales":Decimal(0),"Mano de obra":Decimal(0),"Equipos":Decimal(0),"Transporte":Decimal(0),"Otros":Decimal(0)}
    rows=[]
    for r in data.resources:
        if r.type not in amounts:raise HTTPException(422,"Tipo de recurso no reconocido")
        qty,price=dec(r.quantity),dec(r.unit_price)
        subtotal=qty*price
        amounts[r.type]+=subtotal
        rows.append({"name":r.name,"type":r.type,"unit":r.unit,"quantity":str(qty),
                     "unit_price":str(price),"subtotal":round2(subtotal)})
    indirect,profit=dec(data.indirect_pct),dec(data.profit_pct)
    if indirect>100 or profit>100:raise HTTPException(422,"Porcentaje fuera de rango")
    direct=sum(amounts.values())
    overhead=direct*indirect/100
    utility=(direct+overhead)*profit/100
    return {"code":data.code,"description":data.description,"unit":data.unit,
        "resources":rows,"categories":{k:round2(v) for k,v in amounts.items()},
        "direct_cost":round2(direct),"indirect_cost":round2(overhead),
        "profit":round2(utility),"unit_price":round2(direct+overhead+utility),
        "formula":"Utilidad aplicada sobre costo directo más indirectos; validar según bases",
        "warning":"Valores netos de referencia, sin IVA. Sin precios de mercado automáticos."}
@router.post("/calcular")
def calculate(data:APU):return compute(data)
@router.post("/excel")
def excel(data:APU):
    result=compute(data);wb=Workbook();ws=wb.active;ws.title="APU"
    ws.append(["CIVILTEK | ANÁLISIS DE PRECIO UNITARIO"])
    ws.append(["Partida",data.code,data.description])
    ws.append(["Unidad",data.unit])
    ws.append([])
    ws.append(["Tipo","Recurso","Unidad","Cantidad","Precio unitario","Subtotal"])
    for r in result["resources"]:
        ws.append([r["type"],r["name"],r["unit"],float(r["quantity"]),float(r["unit_price"]),r["subtotal"]])
    ws.append(["COSTO DIRECTO","","","","",result["direct_cost"]])
    ws.append(["GASTOS INDIRECTOS (%)",data.indirect_pct,"","","",result["indirect_cost"]])
    ws.append(["UTILIDAD (%)",data.profit_pct,"","","",result["profit"]])
    ws.append(["PRECIO UNITARIO","","","","",result["unit_price"]])
    for cell in ws[5]:
        cell.fill=PatternFill("solid",fgColor="143949");cell.font=Font(bold=True,color="FFFFFF")
    ws["A1"].font=Font(size=15,bold=True,color="168F88")
    for col,width in {"A":28,"B":52,"C":16,"D":17,"E":20,"F":23}.items():ws.column_dimensions[col].width=width
    for row in ws.iter_rows(min_row=6,min_col=4,max_col=6):
        for cell in row:cell.number_format='#,##0.00;[Red](#,##0.00)'
    ws.freeze_panes="C6"
    output=BytesIO();wb.save(output);output.seek(0)
    return StreamingResponse(output,media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition":'attachment; filename="civiltek-apu.xlsx"'})
